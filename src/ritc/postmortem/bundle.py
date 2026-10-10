"""
One report for everything: `reports/report-<date>-<time>_<local|live>.html`, `.xlsx` and `.json`.

Contains a plain-language explanation of the case, every bot run reviewed (from its log and journal), and with
`--stress` the current bot through every stress scenario. `_live` if any run traded the RIT server.
"""

from __future__ import annotations

import json
import time
from pathlib import Path

from .findings import AB_RESULTS, IMPROVEMENTS
from .html import EXPLAINER, catalogue, overview, page, run_section


def _run_tender_rows(models: list[dict]) -> list[dict]:
    from .simmatch import stream
    from .stress_report import _levels
    cache: dict = {}
    out = []
    for m in models:
        for r in m["ledger"]:
            seed = (r.get("stream") or (None, None))[0]
            if seed is not None and seed not in cache:
                cache[seed] = stream(seed, books_all=True)["books"]
            for d in (r["decisions"] or [None]):
                tick = (d or {}).get("tick") or r.get("seen_tick") or r.get("arrival_tick")
                book = {"tick": tick, "bid": [], "ask": []}
                if seed is not None and tick is not None:
                    b = cache[seed][max(1, min(420, int(tick))) - 1].get(r["ticker"])
                    book.update(bid=_levels(b, "bid"), ask=_levels(b, "ask"))
                mid = (book["bid"][0][0] + book["ask"][0][0]) / 2 if book["bid"] and book["ask"] else None
                ans = "not answered" if d is None else ("ACCEPT" if d["accept"] else "decline")
                out.append({"tid": r["tid"], "kind": r["kind"], "action": r["action"], "ticker": r["ticker"],
                            "qty": r["qty"], "price": d["price"] if d else r.get("price"),
                            "tick": r.get("arrival_tick") or r.get("seen_tick"), "expires": r.get("expires"),
                            "answer_tick": tick if d else None, "answer": ans, "pps": (d or {}).get("pps"),
                            "bid": (d or {}).get("bid"), "mid_answer": mid, "mid": r.get("mid_arrival"),
                            "reserve": r.get("reserve"), "rival": r.get("rival"),
                            "outcome": "rejected" if d and d.get("rejected") else "", "book": book, "seed": seed,
                            "realized": None, "source": f"run {m['stamp']} ({m['mode']})"})
    return out


def write_xlsx(models: list[dict], stress: dict | None, path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    from .xlsx import BLUE, CENTS, GREEN, LEVELS, MONEY, PRICE, _f, _header, _layout, _live_extras, _put
    wb = Workbook()
    runs = wb.active
    runs.title = "Runs"
    _header(runs, 1, ["Run", "Mode", "Bot", "Seed(s)", "Tenders", "Accepts", "Gaps", "Journal P&L",
                      "Current bot on this market"])
    for i, m in enumerate(models, 2):
        rep = m["replays"][0] if m["replays"] else None
        vals = [m["stamp"], m["mode"], m["run"]["version"], ", ".join(map(str, m["market"]["seeds"])) or "-",
                m["run"]["tenders"], m["run"]["accepts"], sum(1 for c in m["checks"] if c["status"] == "gap"),
                (m.get("journal") or {}).get("final_pnl"), rep["score"] if rep else None]
        for j, v in enumerate(vals, 1):
            runs.cell(row=i, column=j, value=v).font = _f(BLUE)
    rows = _run_tender_rows(models)
    if stress:
        from .stress_report import add_stress_sheets, stress_tender_rows
        add_stress_sheets(wb, stress["results"], stress["sums"])
        rows += stress_tender_rows(stress["results"], stress["sums"])
    ten, books, rep_ws, live = (wb.create_sheet(x) for x in ("Tenders", "Books", "Tender replay", "Live (RTD)"))
    _header(ten, 1, ["Row", "Tender ID", "Type", "Side (ours)", "Security", "Volume (+buy)", "Price", "Tick received",
                     "Tick expired", "Answer tick", "Bot answer", "Bot est./share", "Bot bid", "Mid at answer",
                     "Mid at arrival", "Edge vs mid at arrival", "Hidden reserve", "Rival bid", "Bid vs reserve",
                     "Outcome", "Gap flag", "Book tick", "Edge declined ($)", "Seed", "Realized/share", "Source"])
    _header(books, 1, ["Row", "Level", "Bid size", "Bid", "Ask", "Ask size"])
    br = 2
    for i, t in enumerate(rows, 1):
        k = i + 1
        sign = 1 if t["action"] == "BUY" else -1
        vals = [i, t["tid"], t["kind"], t["action"], t["ticker"], sign * t["qty"], t["price"], t["tick"], t["expires"],
                t["answer_tick"], t["answer"], t["pps"], t["bid"], t["mid_answer"], t["mid"], None, t["reserve"],
                t["rival"], None, t["outcome"], None, t["book"]["tick"], None, t["seed"], t["realized"], t["source"]]
        for j, v in enumerate(vals, 1):
            if v is not None:
                ten.cell(row=k, column=j, value=v).font = _f(BLUE)
        ten.cell(row=k, column=16, value=f'=IF(OR(G{k}="",O{k}=""),"",SIGN(F{k})*(O{k}-G{k}))')
        ten.cell(row=k, column=19, value=f'=IF(OR(M{k}="",Q{k}=""),"",SIGN(F{k})*(M{k}-Q{k}))')
        ten.cell(row=k, column=21, value=(
            f'=IF(K{k}="not answered","not answered",IF(AND(K{k}="decline",C{k}="private",N(P{k})>0.1),'
            f'"declined: 10c+ inside the mid",IF(AND(S{k}<>"",N(S{k})<0),"bid short of the reserve","")))'))
        ten.cell(row=k, column=23, value=f'=IF(U{k}="declined: 10c+ inside the mid",P{k}*ABS(F{k}),0)')
        for col, fmt in ((7, PRICE), (12, CENTS), (13, PRICE), (14, PRICE), (15, PRICE), (16, CENTS), (17, PRICE),
                         (18, PRICE), (19, CENTS), (23, MONEY), (25, CENTS)):
            ten.cell(row=k, column=col).number_format = fmt
        for lv in range(LEVELS):
            bb = t["book"]["bid"][lv] if lv < len(t["book"]["bid"]) else (None, None)
            aa = t["book"]["ask"][lv] if lv < len(t["book"]["ask"]) else (None, None)
            for j, v in enumerate([i, lv + 1, bb[1], bb[0], aa[0], aa[1]], 1):
                if v is not None:
                    books.cell(row=br, column=j, value=v)
            br += 1
    ten.freeze_panes = "C2"
    tend, bend = max(2, len(rows) + 1), max(2, br - 1)
    _layout(rep_ws, "data", sel="$B$13", tend=tend, bend=bend, single_book=True)
    _put(rep_ws, "A13", "Show Tenders row (type a number)", bold=True)
    _put(rep_ws, "B13", 1, BLUE, fill=PatternFill("solid", fgColor="FFFF00"), note=f"1 to {len(rows)}: a row of Tenders.")
    _put(rep_ws, "A14", "Source / seed of that row")
    _put(rep_ws, "B14", f'=INDEX(Tenders!$Z$1:$Z${tend},$B$13+1)&" / seed "&INDEX(Tenders!$X$1:$X${tend},$B$13+1)', GREEN)
    _layout(live, "rtd")
    _live_extras(live)
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return path


def write_bundle(models: list[dict], stress: dict | None, out_dir: str = "reports", progress=print) -> list[Path]:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    mode = "live" if any(m["mode"] == "live" for m in models) else "local"
    stamp = time.strftime("%Y%m%d-%H%M%S")
    stem = out / f"report-{stamp}_{mode}"
    body = (f'<h1 id="top">Liability bot report {stamp} ({mode})</h1><p class=sub>Liquidity Risk Case (2027 selection) · '
            f"{len(models)} bot run(s)" + (" · stress test of the current bot" if stress else "")
            + " · the same content is in this report's .xlsx and .json</p>" + EXPLAINER + overview(models)
            + "".join(run_section(m) for m in models))
    if stress:
        from .stress_report import stress_block
        body += stress_block(stress["results"], stress["sums"], stress["stamp"], stress["seeds"])
    body += catalogue(IMPROVEMENTS)
    stem.with_suffix(".html").write_text(page(f"Liability report {stamp}", body, "Runs and stress test"))
    from .live import learnings
    doc = {"generated": stamp, "mode": mode, "improvements": IMPROVEMENTS, "ab_results": AB_RESULTS,
           "runs": [{"stamp": m["stamp"], "mode": m["mode"], "mode_why": m["mode_why"], "run": m["run"],
                     "market": m["market"], "checks": m["checks"],
                     "learnings": learnings(m["stamp"], m["mode"], m["journal"]) if m.get("journal") else None,
                     "replays": [{k: v for k, v in r.items() if k not in ("log", "inventory", "tenders")}
                                 for r in m["replays"]]} for m in models]}
    if stress:
        from .stress_report import combined_json
        doc["stress"] = combined_json(stress["results"], stress["sums"], stress["stamp"], stress["seeds"])
    stem.with_suffix(".json").write_text(json.dumps(doc, default=str))
    written = [stem.with_suffix(".html"), stem.with_suffix(".json")]
    try:
        written.append(write_xlsx(models, stress, stem.with_suffix(".xlsx")))
    except ImportError as exc:
        progress(f"  no .xlsx ({exc}); pip install -e \".[report]\"")
    return written
