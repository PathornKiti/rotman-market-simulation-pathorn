"""
Excel workbook per run, laid out like the case's `Liquidity help file.xlsx`.

Sheets
    Tender replay   the help file's layout (current tick, tender block, 20-level books of CRZY / TAME / CROC)
                    driven by one input cell: pick a row of `Tenders` and the sheet shows that tender and the
                    books at the tick the bot answered, then prices it off the book with formulas
                    (Objective 1: walk the book, commission, profit/share, decision, auction bid)
    Tenders         every tender and decision in the log, with the simulator's truth; edge, auction margin and
                    gap flags are formulas
    Books           the order books behind `Tender replay` (3 stocks x 20 levels per row of `Tenders`)
    Summary         run facts and the gap checks; counts are formulas over `Tenders`
    Current bot     the current bot replayed on the same market: P&L parts and the tenders it took
    Live (RTD)      the help file itself, fixed and extended, for live use with the RIT client:
                    robust TENDERINFO parsing, CROC's security number, ticks left to decide, the same
                    evaluation formulas, the next two open tenders, positions and limit room

Inputs are blue, links to other sheets green, formulas black; yellow cells are the ones to change.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TICKERS = ("CRZY", "TAME", "CROC")
LEVELS = 20
ROWS_PER_SNAP = LEVELS * len(TICKERS)
FONT = "Arial"
BLUE, GREEN, BLACK, GREY = "0000FF", "008000", "000000", "666666"
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEAD = PatternFill("solid", fgColor="E7E6E6")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
PRICE, SHARES, MONEY = "0.00;-0.00;;@", "#,##0;-#,##0;;@", "$#,##0;($#,##0);-"
CENTS = '0.000;-0.000;"-"'
# ladder blocks: (ticker, level col, cum bid, bid size, bid, ask, ask size, cum ask, sec-number cell, ticker cell)
BLOCKS = (("CRZY", "G", "H", "I", "J", "K", "L", "M", "H1", "I1"),
          ("TAME", "O", "P", "Q", "R", "S", "T", "U", "P1", "Q1"),
          ("CROC", "W", "X", "Y", "Z", "AA", "AB", "AC", "X1", "Y1"))


def _f(color=BLACK, bold=False, size=10, italic=False):
    return Font(name=FONT, color=color, bold=bold, size=size, italic=italic)


def _put(ws, ref, value, color=BLACK, fmt=None, bold=False, fill=None, note=None, italic=False, align=None):
    c = ws[ref]
    c.value = value
    c.font = _f(color, bold, italic=italic)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if note:
        c.comment = Comment(note, "ritc report")
    if align:
        c.alignment = Alignment(horizontal=align)
    return c


def _header(ws, row, labels, start_col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=lab)
        c.font = _f(bold=True)
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")


# ------------------------------------------------------------------------------- help-file layout + evaluation
def _layout(ws, source: str, sel: str | None = None, tend: int = 2, bend: int = 61,
            single_book: bool = False) -> None:
    """
    The help file's cells A1:AC21 plus the evaluation block. `source` = "rtd" (live RTD links) or "data"
    (INDEX into the Tenders / Books sheets, row picked by the input cell `sel`). `single_book`: Books holds only
    the tender's own stock (20 rows per tender: Row, Level, Bid size, Bid, Ask, Ask size); the other ladders stay 0.
    """
    ws.column_dimensions["A"].width = 34
    for col in ("B", "C", "D"):
        ws.column_dimensions[col].width = 12
    ws.column_dimensions["E"].width = 30
    for i in range(7, 30):
        ws.column_dimensions[get_column_letter(i)].width = 9
    rtd = source == "rtd"
    T = "Tenders!"
    row = f"{sel}+1" if sel else None

    def tcol(col: str) -> str:            # the selected row of a Tenders column
        return f"INDEX({T}${col}$1:${col}${tend},{row})"

    _put(ws, "A1", "Current tick", bold=True)
    if rtd:
        _put(ws, "B1", '=IFERROR(B2-RTD("rit2.rtd",,"TIMEREMAINING"),"")')
    else:
        _put(ws, "B1", f"={tcol('J')}", GREEN, note="Tick the bot answered this tender (from the log).")
    _put(ws, "A2", "Total tick", bold=True)
    _put(ws, "B2", 420, BLUE, note="Brief: the case runs for 420 seconds.")
    ws.merge_cells("A3:E3")
    _put(ws, "A3", "Tenders", bold=True, fill=HEAD)
    labels = ["Tender ID", "Security", "Sec Number", "Volume", "Price", "Tick received", "Tick expired"]
    for i, lab in enumerate(labels):
        _put(ws, f"A{4 + i}", lab)
    if rtd:
        _put(ws, "D4", "Full text", GREY)
        _put(ws, "E4", '=IFERROR(RTD("rit2.rtd",,"TENDERINFO",1),"")',
             note='RTD doc: "ID,Ticker,Volume,Price,tick received,tick expires", e.g. 11,TAME,58000,24.51,49,79')
        for i in range(1, 6):
            _put(ws, f"D{4 + i}", f"remain {i}", GREY)
            # Robust split: cut at the comma position, not at LEN() of the value already parsed (the help file's
            # RIGHT(E4,LEN(E4)-LEN(B4)-1) breaks when the text is "24.50" but the parsed value prints as 24.5).
            _put(ws, f"E{4 + i}", f'=IFERROR(MID(E{3 + i},FIND(",",E{3 + i})+1,999),"")')
        _put(ws, "B4", '=IFERROR(VALUE(LEFT(E4,FIND(",",E4)-1)),"")')
        _put(ws, "B5", '=IFERROR(LEFT(E5,FIND(",",E5)-1),"")')
        _put(ws, "B7", '=IFERROR(VALUE(LEFT(E6,FIND(",",E6)-1)),"")')
        _put(ws, "B8", '=IFERROR(VALUE(LEFT(E7,FIND(",",E7)-1)),"")',
             note="Auctions may carry no price (or 0): the evaluation then shows the bid to submit.")
        _put(ws, "B9", '=IFERROR(VALUE(LEFT(E8,FIND(",",E8)-1)),"")')
        _put(ws, "B10", '=IFERROR(VALUE(E9),"")')
    else:
        _put(ws, "B4", f"={tcol('B')}", GREEN)
        _put(ws, "B5", f"={tcol('E')}", GREEN)
        _put(ws, "B7", f"={tcol('F')}", GREEN, SHARES, note="+ = we BUY the block (RTD sign convention).")
        _put(ws, "B8", f'=IF({tcol("G")}="","",{tcol("G")})', GREEN, PRICE)
        _put(ws, "B9", f"={tcol('H')}", GREEN)
        _put(ws, "B10", f"={tcol('I')}", GREEN)
        _put(ws, "D4", "Full text", GREY)
        _put(ws, "E4", '=IF(B4="","",B4&","&B5&","&B7&","&B8&","&B9&","&B10)', GREY,
             note="Same format as RTD TENDERINFO.")
    _put(ws, "B6", '=IF(B5="","",IF(B5="CRZY",1,IF(B5="TAME",2,3)))')
    _put(ws, "C7", '=IFERROR(IF(B7>0,"BUY",IF(B7<0,"SELL","")),"")')
    _put(ws, "A11", "Ticks left to decide")
    _put(ws, "B11", '=IF(OR(B10="",B1=""),"",B10-B1)',
         note="Decision window still open. Trading this stock before the answer is front-running (brief).")

    # Books: the help file's three ladders.
    _put(ws, "G1", "Security", bold=True)
    for tk, lv, cb, bs, bp, ap, asz, ca, numcell, tcell in BLOCKS:
        _put(ws, numcell, TICKERS.index(tk) + 1, BLUE)
        _put(ws, tcell, tk, BLUE, bold=True, fmt="@")
        for k in range(1, LEVELS + 1):
            r = k + 1
            _put(ws, f"{lv}{r}", k, GREY)
            _put(ws, f"{cb}{r}", f"=SUM({bs}$2:{bs}{r})", fmt=SHARES)
            _put(ws, f"{ca}{r}", f"=SUM({asz}$2:{asz}{r})", fmt=SHARES)
            if rtd:
                for col, field, fmt in ((bs, "AGBSZ", SHARES), (bp, "AGBID", PRICE), (ap, "AGASK", PRICE),
                                        (asz, "AGASZ", SHARES)):
                    # -- turns a numeric string into a number and an empty level into an error -> 0
                    _put(ws, f"{col}{r}", f'=IFERROR(--RTD("rit2.rtd",,${tcell[0]}${tcell[1:]},"{field}",{lv}{r}),0)',
                         fmt=fmt)
            else:
                if single_book:
                    base = f"2+({sel}-1)*{LEVELS}+{lv}{r}-1"
                    srcs = (("C", bs, SHARES), ("D", bp, PRICE), ("E", ap, PRICE), ("F", asz, SHARES))
                    for src, col, fmt in srcs:
                        _put(ws, f"{col}{r}", f"=IF($B$5=${tcell[0]}${tcell[1:]},INDEX(Books!${src}$1:${src}${bend},"
                                              f"{base}),0)", GREEN, fmt)
                    continue
                base = f"2+({sel}-1)*{ROWS_PER_SNAP}+(${numcell[0]}${numcell[1:]}-1)*{LEVELS}+{lv}{r}-1"
                for col, src, fmt in ((bs, "D", SHARES), (bp, "E", PRICE), (ap, "F", PRICE), (asz, "G", SHARES)):
                    _put(ws, f"{col}{r}", f"=INDEX(Books!${src}$1:${src}${bend},{base})", GREEN, fmt)
    _evaluation(ws, rtd, row, tend)


def _evaluation(ws, rtd: bool, row: str | None, tend: int = 2) -> None:
    """Objective 1 with formulas: walk the book for the block on the side we unwind into."""
    _put(ws, "A23", "Tender evaluation (Objective 1: price it off the order book)", bold=True, fill=HEAD)
    ws.merge_cells("A23:E23")
    inputs = [("A25", "Commission per share (brief: $0.02)", "B25", 0.02, CENTS),
              ("A26", "Book refill factor (bot: 2.0)", "B26", 2, "0.0"),
              ("A27", "Min profit per share (bot: 0.01)", "B27", 0.01, CENTS),
              ("A28", "Exit cost if held to the bell (per share)", "B28", 0.01, CENTS),
              ("A29", "Auction margin (bot: 0.05)", "B29", 0.05, CENTS)]
    notes = {"B26": "The visible book refills while you unwind; the bot walks the book with every level x 2.",
             "B28": "Brief: open positions close at the last traded price with no fine. ~1c covers adverse "
                    "resting fills and the close at the bid/ask.",
             "B29": "The bot bids its value minus 5c on competitive auctions and winner-take-all."}
    for la, lab, ref, val, fmt in inputs:
        _put(ws, la, lab)
        _put(ws, ref, val, BLUE, fmt, fill=YELLOW, note=notes.get(ref))
    rows = [
        ("A31", "Unwind side", "B31",
         '=IF(B7="","",IF(B7>0,"SELL into the bids",IF(B7<0,"BUY from the asks","")))', None),
        ("A32", "Shares to unwind", "B32", '=IF(B7="","",ABS(B7))', SHARES),
        ("A33", "Mid now", "B33", '=IFERROR((CHOOSE(B6,J2,R2,Z2)+CHOOSE(B6,K2,S2,AA2))/2,"")', PRICE),
        ("A34", "Exit VWAP walking the book (x refill)", "B34",
         '=IF(OR(B32="",K46=0),"",(L46+(B32-K46)*(INDEX(H26:H45,COUNTIF(K26:K45,">0"))-SIGN(B7)*0.05))/B32)', "0.0000"),
        ("A35", "Shares beyond the visible book", "B35", '=IF(B32="","",B32-K46)', SHARES),
        ("A36", "Profit/share walking the book, after commission", "B36",
         '=IF(OR(B8="",B8=0,B34=""),"",SIGN(B7)*(B34-B8)-B25)', CENTS),
        ("A37", "Total profit walking the book", "B37", '=IF(B36="","",B36*B32)', MONEY),
        ("A38", "Edge vs the mid per share", "B38", '=IF(OR(B8="",B8=0,B33=""),"",SIGN(B7)*(B33-B8))', CENTS),
        ("A39", "Profit/share if held to the bell", "B39", '=IF(B38="","",B38-B28)', CENTS),
        ("A40", "Decision (walk the book vs min profit)", "B40",
         '=IF(B7="","",IF(OR(B8="",B8=0),"AUCTION: bid "&TEXT(B41,"0.00"),IF(B36>=B27,"ACCEPT","DECLINE")))', None),
        ("A41", "Auction: most aggressive bid that keeps the margin", "B41",
         '=IF(B34="","",ROUND(B34-SIGN(B7)*(B25+B29),2))', PRICE),
    ]
    for la, lab, ref, formula, fmt in rows:
        _put(ws, la, lab)
        _put(ws, ref, formula, fmt=fmt)
    ws["B40"].font = _f(bold=True)
    _put(ws, "A42", "Limits: brief 250,000 gross / 100,000 net shares (enforced)", GREY, italic=True)
    # helper table: walk the book
    _header(ws, 25, ["Level", "Price", "Size x refill", "Before", "Taken", "Value"], start_col=7)
    for k in range(1, LEVELS + 1):
        r, src = 25 + k, k + 1
        _put(ws, f"G{r}", k, GREY)
        _put(ws, f"H{r}", f'=IF($B$7="",0,IF($B$7>0,CHOOSE($B$6,J{src},R{src},Z{src}),'
                          f'CHOOSE($B$6,K{src},S{src},AA{src})))', fmt=PRICE)
        _put(ws, f"I{r}", f'=IF($B$7="",0,IF($B$7>0,CHOOSE($B$6,I{src},Q{src},Y{src}),'
                          f'CHOOSE($B$6,L{src},T{src},AB{src}))*$B$26)', fmt=SHARES)
        _put(ws, f"J{r}", "=0" if k == 1 else f"=J{r - 1}+I{r - 1}", fmt=SHARES)
        _put(ws, f"K{r}", f'=IF($B$32="",0,MIN(I{r},MAX(0,$B$32-J{r})))', fmt=SHARES)
        _put(ws, f"L{r}", f"=H{r}*K{r}", fmt="#,##0.00;-#,##0.00;;@")
    _put(ws, "G46", "Total", bold=True)
    _put(ws, "K46", "=SUM(K26:K45)", fmt=SHARES, bold=True)
    _put(ws, "L46", "=SUM(L26:L45)", fmt="#,##0.00", bold=True)
    _put(ws, "G47", "Shares past the last level are priced at the worst level -5c (as the bot does).", GREY,
         italic=True)
    if not rtd and row:
        T = "Tenders!"

        def blank_or(col: str) -> str:
            ref = f"INDEX({T}${col}$1:${col}${tend},{row})"
            return f'=IF({ref}="","",{ref})'
        extra = [("A44", "What the bot did (from the log)", None, None, None),
                 ("A45", "Bot's answer", "B45", f"=INDEX({T}$K$1:$K${tend},{row})", None),
                 ("A46", "Bot's estimate per share", "B46", blank_or("L"), CENTS),
                 ("A47", "Bot's bid", "B47", blank_or("M"), PRICE),
                 ("A48", "Hidden reserve (simulator)", "B48", blank_or("Q"), PRICE),
                 ("A49", "Rival bid, winner-take-all (simulator)", "B49", blank_or("R"), PRICE),
                 ("A50", "Bot's extra charges vs this sheet (risk, drift, crowd)", "B50",
                  '=IF(OR(B46="",B36=""),"",B36-B46)', CENTS),
                 ("A51", "Gap flag", "B51", f"=INDEX({T}$U$1:$U${tend},{row})", None)]
        for la, lab, ref, formula, fmt in extra:
            _put(ws, la, lab, bold=ref is None, fill=HEAD if ref is None else None)
            if ref:
                _put(ws, ref, formula, GREEN, fmt)
        ws["B51"].font = _f(GREEN, bold=True)


def _live_extras(ws) -> None:
    """Below the evaluation: the next two open tenders, positions and limit room (RTD)."""
    _put(ws, "A54", "Other open tenders (RTD TENDERINFO 2 and 3)", bold=True, fill=HEAD)
    ws.merge_cells("A54:E54")
    _header(ws, 55, ["No.", "Full text", "ID", "Security", "Volume", "Price", "Received", "Expires"])
    for i, n in enumerate((2, 3)):
        r = 56 + i
        _put(ws, f"A{r}", n, GREY)
        _put(ws, f"B{r}", f'=IFERROR(RTD("rit2.rtd",,"TENDERINFO",{n}),"")')
        # parse with SUBSTITUTE-based field extraction (field k of a comma list)
        for j, col in enumerate("CDEFGH"):
            _put(ws, f"{col}{r}", f'=IFERROR(TRIM(MID(SUBSTITUTE($B{r},",",REPT(" ",99)),{j}*99+1,99)),"")')
    _put(ws, "A59", "Positions and limits", bold=True, fill=HEAD)
    ws.merge_cells("A59:E59")
    _header(ws, 60, ["Stock", "Position", "If this tender is accepted"])
    for i, tk in enumerate(TICKERS):
        r = 61 + i
        _put(ws, f"A{r}", tk, BLUE)
        _put(ws, f"B{r}", f'=IFERROR(--RTD("rit2.rtd",,"{tk}","POSITION"),0)', fmt=SHARES)
        _put(ws, f"C{r}", f'=B{r}+IF($B$5=A{r},$B$7,0)', fmt=SHARES)
    rows = [("A64", "Gross (sum of |positions|)", "B64", "=ABS(B61)+ABS(B62)+ABS(B63)", "C64",
             "=ABS(C61)+ABS(C62)+ABS(C63)"),
            ("A65", "Net (sum of positions)", "B65", "=B61+B62+B63", "C65", "=C61+C62+C63")]
    for la, lab, rb, fb, rc, fc in rows:
        _put(ws, la, lab)
        _put(ws, rb, fb, fmt=SHARES)
        _put(ws, rc, fc, fmt=SHARES)
    _put(ws, "A66", "Gross limit / net limit (brief)")
    _put(ws, "B66", 250000, BLUE, SHARES, fill=YELLOW)
    _put(ws, "C66", 100000, BLUE, SHARES, fill=YELLOW)
    _put(ws, "A67", "Fits the limits if accepted?")
    _put(ws, "B67", '=IF(B7="","",IF(AND(C64<=B66,ABS(C65)<=C66),"yes","NO - the server will refuse it"))')
    ws["B67"].font = _f(bold=True)
    _put(ws, "A68", "Trader P&L (RTD)")
    _put(ws, "B68", '=IFERROR(RTD("rit2.rtd",,"PL"),"")', fmt=MONEY)
    _put(ws, "A70", "Speculation fine for opening N new shares (brief)", bold=True, fill=HEAD)
    ws.merge_cells("A70:E70")
    _put(ws, "A71", "Shares that open a new position")
    _put(ws, "B71", 0, BLUE, SHARES, fill=YELLOW)
    _put(ws, "A72", "Fine ($0.20 first 5,000, $0.40 beyond)")
    _put(ws, "B72", "=0.2*MIN(B71,5000)+0.4*MAX(0,B71-5000)", fmt=MONEY)


# ------------------------------------------------------------------------------------------------- data
def _book_levels(book: dict | None, side: str) -> list[tuple[float, int]]:
    if not book:
        return []
    agg: dict[float, int] = {}
    for o in book.get(side, []):
        p = o.get("price")
        if p is None:
            continue
        agg[float(p)] = agg.get(float(p), 0) + int(o.get("quantity", 0)) - int(o.get("quantity_filled", 0) or 0)
    return sorted(agg.items(), reverse=(side == "bid"))


def _tender_rows(m: dict) -> list[dict]:
    out = []
    for r in m["ledger"]:
        decs = r["decisions"] or [None]
        for d in decs:
            out.append({"r": r, "d": d})
    return out


def write_xlsx(m: dict, path: Path) -> Path:
    from .simmatch import stream
    wb = Workbook()
    rep_ws = wb.active
    rep_ws.title = "Tender replay"
    ten = wb.create_sheet("Tenders")
    books = wb.create_sheet("Books")
    summ = wb.create_sheet("Summary")
    cur = wb.create_sheet("Current bot")
    live = wb.create_sheet("Live (RTD)")

    rows = _tender_rows(m)
    tend, bend = max(2, len(rows) + 1), 1 + max(1, len(rows)) * ROWS_PER_SNAP    # never an empty range
    seeds = m["market"]["seeds"]
    st_cache = {s: stream(s, books_all=True) for s in seeds}

    # ---------------- Tenders
    heads = ["Row", "Tender ID", "Type", "Side (ours)", "Security", "Volume (+buy)", "Price", "Tick received",
             "Tick expired", "Answer tick", "Bot answer", "Bot est./share", "Bot bid", "Mid at answer",
             "Mid at arrival", "Edge vs mid at arrival", "Hidden reserve", "Rival bid", "Bid vs reserve", "Market",
             "Gap flag", "Book tick", "Edge declined ($)"]
    _header(ten, 1, heads)
    widths = [5, 9, 9, 8, 9, 11, 8, 9, 9, 9, 13, 10, 8, 9, 9, 10, 9, 9, 9, 22, 30, 8, 11]
    for i, w in enumerate(widths, 1):
        ten.column_dimensions[get_column_letter(i)].width = w
    ten.freeze_panes = "C2"
    for i, x in enumerate(rows, 1):
        r, d = x["r"], x["d"]
        k = i + 1
        sign = 1 if r["action"] == "BUY" else -1
        ans = "not answered" if d is None else ("ACCEPT" if d["accept"] else "decline")
        if d and d.get("rejected"):
            ans = "ACCEPT (rejected)"
        tick = (d or {}).get("tick") or r.get("seen_tick") or r.get("arrival_tick")
        seed = (r.get("stream") or (None, None))[0]
        bt = None
        if seed in st_cache and tick is not None:
            bt = max(1, min(420, int(tick)))
        vals = [i, r["tid"], r["kind"], r["action"], r["ticker"], sign * r["qty"],
                d["price"] if d else r.get("price"), r.get("arrival_tick") or r.get("seen_tick"),
                r.get("expires"), tick, ans, (d or {}).get("pps"), (d or {}).get("bid"), None, r.get("mid_arrival"),
                None, r.get("reserve"), r.get("rival"), None, (d or {}).get("variant") or "", None, bt]
        for j, v in enumerate(vals, 1):
            c = ten.cell(row=k, column=j, value=v)
            c.font = _f(BLUE)
        at = f'2+(A{k}-1)*{ROWS_PER_SNAP}+(MATCH(E{k},{{"CRZY","TAME","CROC"}},0)-1)*{LEVELS}'
        bE, bF = f"Books!$E$1:$E${bend}", f"Books!$F$1:$F${bend}"
        ten.cell(row=k, column=14, value=f'=IFERROR(IF(INDEX({bE},{at})=0,"",(INDEX({bE},{at})'
                                          f'+INDEX({bF},{at}))/2),"")').font = _f(GREEN)
        ten.cell(row=k, column=16, value=f'=IF(OR(G{k}="",O{k}=""),"",SIGN(F{k})*(O{k}-G{k}))').font = _f()
        ten.cell(row=k, column=19, value=f'=IF(OR(M{k}="",Q{k}=""),"",SIGN(F{k})*(M{k}-Q{k}))').font = _f()
        ten.cell(row=k, column=21, value=(
            f'=IF(K{k}="not answered","not answered",IF(AND(K{k}="decline",P{k}<>"",N(P{k})>0.1),'
            f'"declined: 10c+ inside the mid",IF(AND(LEFT(K{k},6)="ACCEPT",P{k}<>"",N(P{k})<0),'
            f'"accepted below the mid",'
            f'IF(AND(S{k}<>"",N(S{k})<0),"bid short of the reserve",""))))')).font = _f()
        ten.cell(row=k, column=23, value=f'=IF(U{k}="declined: 10c+ inside the mid",P{k}*ABS(F{k}),0)').font = _f()
        for col, fmt in ((6, SHARES), (7, PRICE), (12, CENTS), (13, PRICE), (14, PRICE), (15, PRICE), (16, CENTS),
                         (17, PRICE), (18, PRICE), (19, CENTS), (23, MONEY)):
            ten.cell(row=k, column=col).number_format = fmt
    n_rows = len(rows)
    last = max(2, n_rows + 1)                      # last data row of Tenders for the summary formulas

    # ---------------- Books (one snapshot per Tenders row, 3 stocks x 20 levels)
    _header(books, 1, ["Row", "Security", "Level", "Bid size", "Bid", "Ask", "Ask size", "Book tick"])
    for c, w in zip("ABCDEFGH", (6, 9, 6, 9, 8, 8, 9, 9)):
        books.column_dimensions[c].width = w
    rr = 2
    for i, x in enumerate(rows, 1):
        r, d = x["r"], x["d"]
        seed = (r.get("stream") or (None, None))[0]
        tick = (d or {}).get("tick") or r.get("seen_tick") or r.get("arrival_tick")
        snap = None
        if seed in st_cache and tick is not None:
            snap = st_cache[seed]["books"][max(1, min(420, int(tick))) - 1]
        for tk in TICKERS:
            bids = _book_levels(snap.get(tk) if snap else None, "bid")
            asks = _book_levels(snap.get(tk) if snap else None, "ask")
            for lv in range(LEVELS):
                b = bids[lv] if lv < len(bids) else (None, None)
                a = asks[lv] if lv < len(asks) else (None, None)
                for j, v in enumerate([i, tk, lv + 1, b[1], b[0], a[0], a[1], tick if snap else None], 1):
                    c = books.cell(row=rr, column=j, value=v)
                    c.font = _f(BLUE if j > 3 else GREY)
                    if j in (5, 6):
                        c.number_format = PRICE
                    elif j in (4, 7):
                        c.number_format = SHARES
                rr += 1
    books.freeze_panes = "A2"

    # ---------------- Tender replay (the help file's layout, driven by B13)
    _layout(rep_ws, "data", sel="$B$13", tend=tend, bend=bend)
    _put(rep_ws, "A13", "Show Tenders row (type a number)", bold=True)
    _put(rep_ws, "B13", 1, BLUE, fill=YELLOW, note=(f"1 to {n_rows}: the row number in the Tenders sheet."
                                                    if n_rows else "This run has no tenders."))
    _put(rep_ws, "C13", f'="of "&{n_rows}', GREY)
    _put(rep_ws, "A14", "Book shown is from tick")
    _put(rep_ws, "B14", f'=IF(INDEX(Tenders!$V$1:$V${tend},$B$13+1)="","no book data",'
                        f'INDEX(Tenders!$V$1:$V${tend},$B$13+1))', GREEN)
    _put(rep_ws, "A15", "Source", GREY)
    src = (f"simulator seed {', '.join(map(str, seeds))}, market without the bot (exact for dry runs; a live bot's "
           "own trades move the real book)" if seeds else "no book data for live runs (record the heat with "
           "`python -m ritc record` to get books)")
    _put(rep_ws, "B15", src, GREY, italic=True)
    rep_ws.freeze_panes = "A2"

    # ---------------- Summary
    run = m["run"]
    summ.column_dimensions["A"].width = 30
    summ.column_dimensions["B"].width = 14
    summ.column_dimensions["C"].width = 110
    _put(summ, "A1", f"Liability run {m['stamp']} ({m['mode']})", bold=True)
    facts = [("Log file", m["file"]), ("Mode", f"{m['mode']}: {m['mode_why']}"), ("Bot version", run["version_label"]),
             ("Bots in the file", run["bots"]), ("Orders", "dry run" if run["dry_run"] else "live"),
             ("Simulator seed(s)", ", ".join(map(str, seeds)) or "-"),
             ("Market variant(s)", "; ".join(f"seed {k}: " + ", ".join(v["scenario"] for v in vs)
                                             for k, vs in m["market"]["variants"].items()) or "-"),
             ("Generated", m["generated"])]
    for i, (k, v) in enumerate(facts, 3):
        _put(summ, f"A{i}", k)
        _put(summ, f"C{i}", v, BLUE)
    r0 = 3 + len(facts) + 1
    _put(summ, f"A{r0}", "Counts (formulas over Tenders)", bold=True, fill=HEAD)
    counts = [("Decision rows", f"=COUNTA(Tenders!A2:A{last})"),
              ("ACCEPT answers", f'=COUNTIF(Tenders!K2:K{last},"ACCEPT*")'),
              ("decline answers", f'=COUNTIF(Tenders!K2:K{last},"decline")'),
              ("not answered", f'=COUNTIF(Tenders!K2:K{last},"not answered")'),
              ("Declined while 10c+ inside the mid", f'=COUNTIF(Tenders!U2:U{last},"declined*")'),
              ("Auction bids short of the reserve", f'=COUNTIF(Tenders!U2:U{last},"bid short*")'),
              ("Edge declined, $ (10c+ cases)", f"=SUM(Tenders!W2:W{last})")]
    for i, (k, f) in enumerate(counts, r0 + 1):
        _put(summ, f"A{i}", k)
        _put(summ, f"B{i}", f, fmt=MONEY if "$" in k else "#,##0")
    r1 = r0 + len(counts) + 2
    _put(summ, f"A{r1}", "Gap checks", bold=True, fill=HEAD)
    _header(summ, r1 + 1, ["Check", "Status", "Finding"])
    for i, c in enumerate(m["checks"], r1 + 2):
        _put(summ, f"A{i}", c["title"])
        _put(summ, f"B{i}", c["status"].upper(), bold=True)
        cell = _put(summ, f"C{i}", c["detail"] + ("\n" + "\n".join(c["evidence"][:6]) if c["evidence"] else ""))
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    r2 = r1 + 3 + len(m["checks"])
    _put(summ, f"A{r2}", "Improvements (trading, not logging)", bold=True, fill=HEAD)
    ids = []
    for c in m["checks"]:
        ids += [i for i in c["improvements"] if i not in ids]
    for i, key in enumerate(ids, r2 + 1):
        imp = m["improvements"][key]
        _put(summ, f"A{i}", imp["title"])
        _put(summ, f"B{i}", imp["status"])
        cell = _put(summ, f"C{i}", imp["what"] + (" Evidence: " + imp["evidence"] if imp.get("evidence") else ""))
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    # ---------------- Current bot
    cur.column_dimensions["A"].width = 34
    for c in "BCDEFGHIJKL":
        cur.column_dimensions[c].width = 12
    if m["replays"]:
        rep = m["replays"][0]
        _put(cur, "A1", f"Current bot replayed on seed {rep['seed']} ({rep['scenario']})", bold=True)
        _put(cur, "A2", "Simulator output (values); totals are formulas.", GREY, italic=True)
        _header(cur, 4, ["P&L part", "$"])
        labels = {"edge": "Tender edge vs mid", "spread_maker": "Resting fills vs mid",
                  "spread_taker": "Crossing fills vs mid",
                  "fees": "Commissions", "impact": "Own price impact", "inventory": "Inventory moves (luck / crowd)",
                  "closeout": "Close-out vs mid", "fines": "Fines"}
        parts = m["parts"]
        for i, k in enumerate(parts, 5):
            _put(cur, f"A{i}", labels[k])
            _put(cur, f"B{i}", round(rep["decomposition"]["total"][k], 2), BLUE, MONEY)
        tr = 5 + len(parts)
        _put(cur, f"A{tr}", "Score (sum of the parts)", bold=True)
        _put(cur, f"B{tr}", f"=SUM(B5:B{tr - 1})", fmt=MONEY, bold=True)
        _put(cur, f"A{tr + 1}", "Score reported by the simulator")
        _put(cur, f"B{tr + 1}", round(rep["score"], 2), BLUE, MONEY)
        _put(cur, f"A{tr + 2}", "Check (should be 0)")
        _put(cur, f"B{tr + 2}", f"=ROUND(B{tr}-B{tr + 1},2)", fmt="0.00")
        h = tr + 4
        _header(cur, h, ["Tender", "Type", "Side", "Security", "Volume", "Price", "Est./share", "Realized $",
                         "Realized/share", "Out resting", "Netted by tender", "Closed at bell"])
        k = h + 1
        for t in rep["tenders"]:
            if t.get("booked_tick") is None:
                continue
            vals = [t["tid"], t["kind"], t["action"], t["ticker"], t["qty"], t["booked_price"], t.get("pps_est"),
                    round(t.get("pnl", 0.0), 2), None, t.get("via_maker", 0), t.get("via_tender", 0),
                    t.get("via_close", 0)]
            for j, v in enumerate(vals, 1):
                c = cur.cell(row=k, column=j, value=v)
                c.font = _f(BLUE)
            cur.cell(row=k, column=9, value=f'=IF(E{k}=0,"",H{k}/E{k})').font = _f()
            for col, fmt in ((5, SHARES), (6, PRICE), (7, CENTS), (8, MONEY), (9, CENTS), (10, SHARES), (11, SHARES),
                             (12, SHARES)):
                cur.cell(row=k, column=col).number_format = fmt
            k += 1
        _put(cur, f"A{k}", "Total", bold=True)
        _put(cur, f"H{k}", f"=SUM(H{h + 1}:H{k - 1})", fmt=MONEY, bold=True)
    else:
        _put(cur, "A1", "No replay (live run: no simulator seed).", GREY)

    # ---------------- Journal (P&L path and struggles), when the run wrote one
    j = m.get("journal")
    if j:
        js = wb.create_sheet("Journal")
        tks = j["tickers"]
        _header(js, 1, ["Tick", "NLV"] + [f"{t} position" for t in tks] + ["Gross", "Net"])
        for i, row in enumerate(j["series"], 2):
            vals = [row["tick"], row["nlv"]] + [row["pos"].get(t, 0) for t in tks]
            for c, v in enumerate(vals, 1):
                js.cell(row=i, column=c, value=v).font = _f(BLUE)
            first, lastc = get_column_letter(3), get_column_letter(2 + len(tks))
            js.cell(row=i, column=3 + len(tks),
                    value="=" + "+".join(f"ABS({get_column_letter(3 + k)}{i})" for k in range(len(tks))) if tks else 0)
            js.cell(row=i, column=4 + len(tks), value=f"=SUM({first}{i}:{lastc}{i})" if tks else 0)
            js.cell(row=i, column=2).number_format = MONEY
        sw = wb.create_sheet("Struggles")
        _header(sw, 1, ["Struggle", "Severity", "Count", "$ at stake", "Detail", "Look at"])
        for i, (k, v) in enumerate(sorted(j["struggles"].items()), 2):
            for c, val in enumerate([k, v["severity"], v["count"], v["value"], v["detail"], v["hint"]], 1):
                sw.cell(row=i, column=c, value=val).font = _f(BLUE)
            sw.cell(row=i, column=4).number_format = MONEY
        n = len(j["struggles"]) + 1
        sw.cell(row=n + 2, column=1, value="Total events").font = _f(bold=True)
        sw.cell(row=n + 2, column=3, value=f"=SUM(C2:C{max(2, n)})")
        for c, w in zip("ABCDEF", (30, 10, 8, 12, 60, 60)):
            sw.column_dimensions[c].width = w

    # ---------------- Live (RTD)
    _layout(live, "rtd")
    _put(live, "A13", "Live sheet: open the RIT client first, then this file (RTD doc).", GREY, italic=True)
    _put(live, "A14", "Fixes vs the help file: robust TENDERINFO split; CROC = 3; ticks left; evaluation below.", GREY,
         italic=True)
    _live_extras(live)

    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 100
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return path
